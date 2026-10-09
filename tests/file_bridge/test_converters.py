import io

from src.services.file_bridge import converters


def _pdf_with_text(text: str) -> bytes:
    """A minimal one-page PDF with ``text`` on it."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_csv_becomes_markdown_table():
    text = converters.convert_csv(b"name,amount\nAlice,10\nBob,20\n")
    assert "Rows: 2 (plus header), columns: 2" in text
    assert "| name | amount |" in text
    assert "| Bob | 20 |" in text


def test_csv_sniffs_semicolons_and_escapes_pipes():
    text = converters.convert_csv(b"a;b\nx|y;2\n")
    assert "| x\\|y | 2 |" in text


def test_xlsx_lists_every_sheet():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Sales"
    workbook.active.append(["region", "total"])
    workbook.active.append(["EU", 42])
    workbook.create_sheet("Empty")
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = converters.convert_xlsx(buffer.getvalue())
    assert "## Sheet: Sales" in text
    assert "| EU | 42 |" in text
    assert "## Sheet: Empty" in text


def test_docx_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_paragraph("Quarterly plan")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "goal", "owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "ship", "Sam"
    buffer = io.BytesIO()
    document.save(buffer)

    text = converters.convert_docx(buffer.getvalue())
    assert "Quarterly plan" in text
    assert "| ship | Sam |" in text


def test_pptx_slides_and_notes():
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.notes_slide.notes_text_frame.text = "Say hello"
    buffer = io.BytesIO()
    presentation.save(buffer)

    text = converters.convert_pptx(buffer.getvalue())
    assert "## Slide 1" in text
    assert "Roadmap" in text
    assert "Notes: Say hello" in text


def test_pdf_text_and_scanned_detection():
    text, looks_scanned = converters.convert_pdf(_pdf_with_text("Invoice total is 42 dollars"))
    assert "Invoice total is 42 dollars" in text
    assert looks_scanned is False

    _, blank_scanned = converters.convert_pdf(_pdf_with_text(""))
    assert blank_scanned is True


def test_text_decoding_handles_bom():
    assert converters.convert_text("﻿hello".encode()) == "hello"
