"""Turn file bytes into plain text, without calling any model.

Each converter returns the full text (up to MAX_TEXT_CHARS); the service layer
decides how much of it goes into the prompt.
"""

import csv
import io
import re

MAX_TEXT_CHARS = 500_000
# A PDF with less text than this per page is treated as scanned.
SCANNED_PDF_CHARS_PER_PAGE = 10


def _cap(text: str) -> str:
    return text[:MAX_TEXT_CHARS]


def _cell(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).replace("|", "\\|").strip()


def _markdown_table(rows: list[list]) -> str:
    if not rows:
        return "(empty)"
    width = max(len(row) for row in rows)
    padded = [[_cell(v) for v in row] + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(row) + " |" for row in padded[1:]]
    return "\n".join(lines)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def convert_text(data: bytes) -> str:
    return _cap(decode_text(data))


def convert_csv(data: bytes, delimiter: str | None = None) -> str:
    text = decode_text(data)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if any(cell.strip() for cell in row)]
    if not rows:
        return "(empty CSV)"
    header = f"Rows: {len(rows) - 1} (plus header), columns: {len(rows[0])}\n\n"
    return _cap(header + _markdown_table(rows))


def convert_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]
            parts.append(
                f"## Sheet: {sheet.title}\nRows: {max(len(rows) - 1, 0)} (plus header)\n\n{_markdown_table(rows)}"
            )
            if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                break
    finally:
        workbook.close()
    return _cap("\n\n".join(parts))


def convert_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parts.append(f"\nTable {index}:\n{_markdown_table(rows)}")
    return _cap("\n".join(parts))


def convert_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                lines.append(_markdown_table(rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append(f"Notes: {slide.notes_slide.notes_text_frame.text.strip()}")
        parts.append(f"## Slide {number}\n" + "\n".join(lines))
    return _cap("\n\n".join(parts))


def convert_pdf(data: bytes) -> tuple[str, bool]:
    """Return (text, looks_scanned)."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        pages.append(f"--- Page {number} ---\n{page_text.strip()}")
    text_chars = sum(len(p) for p in pages) - sum(len(f"--- Page {n} ---\n") for n in range(1, len(pages) + 1))
    looks_scanned = len(pages) > 0 and text_chars < SCANNED_PDF_CHARS_PER_PAGE * len(pages)
    return _cap("\n\n".join(pages)), looks_scanned
